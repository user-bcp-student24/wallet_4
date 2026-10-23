from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "Wallet_para_pagos_cotidianos.pptx"

GREEN = "174B3D"
DARK = "12342E"
LIME = "D6F58B"
PAPER = "F6F8F5"
INK = "202B27"
MUTED = "68756E"
LINE = "DCE2DC"
CORAL = "D76C52"
WHITE = "FFFFFF"
PALE_GREEN = "E8F0E7"
PALE_CORAL = "F8EAE5"
PALE_YELLOW = "FBF2D9"
FONT = "Aptos"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def color(value):
    return RGBColor.from_string(value)


def textbox(slide, text, x, y, w, h, size=18, fill=INK, bold=False,
            align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, font=FONT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.02)
    frame.margin_right = Inches(0.02)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color(fill)
    return shape


def panel(slide, x, y, w, h, fill=WHITE, line=LINE, rounded=True):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    shape.line.color.rgb = color(line)
    shape.line.width = Pt(0.8)
    return shape


def bg(slide, fill=PAPER):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color(fill)


def header(slide, eyebrow, title, subtitle=None, dark=False):
    base = WHITE if dark else INK
    accent = LIME if dark else GREEN
    textbox(slide, eyebrow.upper(), 0.72, 0.42, 11.9, 0.28, 10, accent, True)
    textbox(slide, title, 0.72, 0.82, 11.9, 0.62, 28, base, True)
    if subtitle:
        textbox(slide, subtitle, 0.74, 1.5, 11.7, 0.5, 14, "DCE9E1" if dark else MUTED)


def footer(slide, number, dark=False):
    c = "A9B9B0" if dark else MUTED
    textbox(slide, "WALLET MVP · PERÚ", 0.72, 7.12, 4.0, 0.18, 9, c, True)
    textbox(slide, f"{number:02d}", 12.0, 7.1, 0.6, 0.2, 9, c, True, PP_ALIGN.RIGHT)


def add_image_contain(slide, path, x, y, w, h, pad=0.12):
    path = Path(path)
    with Image.open(path) as image:
        iw, ih = image.size
    max_w, max_h = w - pad * 2, h - pad * 2
    scale = min(max_w / iw, max_h / ih)
    out_w, out_h = iw * scale, ih * scale
    left = x + (w - out_w) / 2
    top = y + (h - out_h) / 2
    slide.shapes.add_picture(str(path), Inches(left), Inches(top),
                             width=Inches(out_w), height=Inches(out_h))


def new_slide(fill=PAPER):
    slide = prs.slides.add_slide(BLANK)
    bg(slide, fill)
    return slide


# 1. Cover
slide = new_slide(DARK)
textbox(slide, "WALLET MVP  ·  PERÚ", 0.8, 0.68, 5.3, 0.3, 11, LIME, True)
textbox(slide, "Pagos cotidianos,\ncon un inicio más claro", 0.8, 1.35, 5.9, 1.8, 34, WHITE, True)
textbox(slide, "W1 propone un onboarding claro.\nLa demo actual muestra el dashboard Angular.", 0.82, 3.42, 5.65, 0.85, 18, "DCE9E1")
panel(slide, 0.82, 5.15, 5.15, 0.78, GREEN, GREEN)
textbox(slide, "PROPUESTA DE PRODUCTO  ·  PROTOTIPO DEMOSTRATIVO", 1.05, 5.4, 4.7, 0.24, 10, LIME, True)
panel(slide, 7.05, 0.82, 5.45, 5.8, WHITE, WHITE)
add_image_contain(slide, ROOT / "frontend-onboarding.png", 7.15, 0.92, 5.25, 5.6, 0.08)
footer(slide, 1, True)

# 2. Need and value proposition
slide = new_slide()
header(slide, "La oportunidad", "La confianza empieza antes del primer pago",
       "El registro debe ayudar a la persona a entender qué pasa y qué sigue.")
textbox(slide, "Una wallet puede perder a una persona antes de su primer pago si el alta genera dudas.",
        0.82, 2.15, 4.15, 1.4, 22, GREEN, True)
items = [
    ("01", "Entender", "Qué información se necesita y en qué paso está."),
    ("02", "Corregir", "Resolver datos incompletos o inválidos sin adivinar."),
    ("03", "Continuar", "Saber cuándo la cuenta está habilitada para pagar."),
]
for i, (num, title, copy) in enumerate(items):
    x = 5.2 + i * 2.55
    panel(slide, x, 2.1, 2.25, 2.75, WHITE, LINE)
    textbox(slide, num, x + 0.22, 2.34, 0.6, 0.35, 12, CORAL, True)
    textbox(slide, title, x + 0.22, 2.94, 1.85, 0.42, 20, GREEN, True)
    textbox(slide, copy, x + 0.22, 3.53, 1.8, 1.0, 13, MUTED)
panel(slide, 0.82, 5.35, 11.95, 0.72, PALE_GREEN, PALE_GREEN)
textbox(slide, "PROPUESTA  ·  Guiar el onboarding, comunicar estados y habilitar pagos solo cuando corresponda.",
        1.08, 5.57, 11.4, 0.3, 15, DARK, True)
footer(slide, 2)

# 3. KYC-aware onboarding
slide = new_slide()
header(slide, "Experiencia objetivo · onboarding", "La verificación define el siguiente paso",
       "La cuenta solo habilita pagos cuando la verificación queda activa.")
panel(slide, 0.82, 2.25, 2.35, 2.0, DARK, DARK)
textbox(slide, "1", 1.08, 2.52, 0.5, 0.4, 20, LIME, True)
textbox(slide, "Registro", 1.08, 3.03, 1.8, 0.38, 20, WHITE, True)
textbox(slide, "Datos y validaciones claras", 1.08, 3.52, 1.85, 0.48, 12, "DCE9E1")
textbox(slide, ">", 3.35, 2.93, 0.45, 0.5, 26, CORAL, True, PP_ALIGN.CENTER)
panel(slide, 3.9, 2.25, 2.55, 2.0, GREEN, GREEN)
textbox(slide, "2", 4.18, 2.52, 0.5, 0.4, 20, LIME, True)
textbox(slide, "Verificación KYC", 4.18, 3.03, 2.05, 0.42, 18, WHITE, True)
textbox(slide, "Resultado y estado de cuenta", 4.18, 3.52, 2.0, 0.48, 12, "DCE9E1")

states = [
    ("ACTIVE", "Wallet habilitada", "Acceso al flujo de pagos", PALE_GREEN, GREEN),
    ("PENDING", "Revisión pendiente", "Informar estado y siguiente paso", PALE_YELLOW, "8A6A18"),
    ("REJECTED", "No habilitada", "Sin acceso a pagos", PALE_CORAL, CORAL),
]
for i, (status, title, copy, fill, accent) in enumerate(states):
    x = 6.85 + i * 1.95
    panel(slide, x, 2.25, 1.72, 2.0, fill, fill)
    textbox(slide, status, x + 0.14, 2.52, 1.45, 0.3, 11, accent, True)
    textbox(slide, title, x + 0.14, 3.0, 1.45, 0.58, 15, INK, True)
    textbox(slide, copy, x + 0.14, 3.65, 1.45, 0.42, 10, MUTED)
panel(slide, 0.82, 5.0, 11.95, 0.9, PALE_GREEN, PALE_GREEN)
textbox(slide, "REGLA DE CONFIANZA", 1.08, 5.2, 2.3, 0.25, 10, GREEN, True)
textbox(slide, "La wallet solo habilita pagos cuando la cuenta está ACTIVE.", 3.45, 5.15, 8.7, 0.38, 18, DARK, True)
footer(slide, 3)

# 4. W1 onboarding journey
slide = new_slide()
header(slide, "Experiencia objetivo · W1", "El onboarding debe resolver una duda por etapa",
       "Cada paso reduce incertidumbre antes de habilitar la wallet.")
steps = [
    ("01", "Registrarse", "Capturar datos mínimos"),
    ("02", "Validar", "Explicar qué corregir"),
    ("03", "Verificar", "Mostrar estado KYC"),
    ("04", "Orientar", "Indicar el siguiente paso"),
    ("05", "Habilitar", "Acceso cuando ACTIVE"),
]
for i, (num, title, copy) in enumerate(steps):
    x = 0.72 + i * 2.53
    panel(slide, x, 2.45, 2.22, 2.0, WHITE, LINE)
    textbox(slide, num, x + 0.18, 2.68, 0.55, 0.28, 11, CORAL, True)
    textbox(slide, title, x + 0.18, 3.14, 1.8, 0.42, 18, GREEN, True)
    textbox(slide, copy, x + 0.18, 3.7, 1.85, 0.55, 11, MUTED)
    if i < 4:
        textbox(slide, ">", x + 2.24, 3.15, 0.28, 0.38, 19, CORAL, True, PP_ALIGN.CENTER)
panel(slide, 0.82, 5.0, 5.75, 0.86, PALE_CORAL, PALE_CORAL)
textbox(slide, "Pendiente de construir", 1.08, 5.16, 2.5, 0.28, 14, CORAL, True)
textbox(slide, "Integrar registro y KYC a Angular.", 3.4, 5.16, 2.9, 0.28, 12, INK)
panel(slide, 6.78, 5.0, 5.99, 0.86, PALE_GREEN, PALE_GREEN)
textbox(slide, "Fuera del alcance W1", 7.04, 5.16, 2.3, 0.28, 14, GREEN, True)
textbox(slide, "Pagos por contacto y QR.", 9.25, 5.16, 3.2, 0.28, 12, INK)
footer(slide, 4)

# 5. Initial and logged-in Angular screens
slide = new_slide()
header(slide, "Demo Angular", "Del registro de prueba al dashboard",
       "Dos estados de una sola aplicación; el cambio de pantalla ocurre localmente.")
panel(slide, 0.72, 2.05, 5.86, 4.7, WHITE, LINE)
add_image_contain(slide, ROOT / "frontend-onboarding.png", 0.84, 2.15, 5.62, 4.48, 0.05)
panel(slide, 6.76, 2.05, 5.86, 4.7, WHITE, LINE)
add_image_contain(slide, ROOT / "frontend-dashboard.png", 6.88, 2.15, 5.62, 4.48, 0.05)
textbox(slide, "INICIO · REGISTRO FICTICIO", 0.9, 6.83, 5.5, 0.24, 10, GREEN, True)
textbox(slide, "DESPUÉS · DASHBOARD DE EJEMPLO", 6.94, 6.83, 5.5, 0.24, 10, GREEN, True)
footer(slide, 5)

# 6. Current demo vs target product
slide = new_slide()
header(slide, "Alcance", "Separar lo que vemos hoy de lo que propone el MVP",
       "La demo visual no debe confundirse con una integración de producción.")
panel(slide, 0.82, 2.18, 5.55, 3.45, DARK, DARK)
textbox(slide, "HOY · DEMOSTRACIÓN", 1.12, 2.48, 4.8, 0.3, 11, LIME, True)
textbox(slide, "• Registro inicial dentro de Angular\n• Validaciones locales del formulario\n• Dashboard con saldos y movimientos ficticios\n• Backend Quarkus: health en 8089",
    1.12, 3.08, 4.8, 1.95, 17, WHITE)
textbox(slide, "No crea cuentas, no valida KYC ni procesa pagos.", 1.12, 5.05, 4.8, 0.36, 12, "DCE9E1", True)
panel(slide, 6.75, 2.18, 5.75, 3.45, WHITE, LINE)
textbox(slide, "SIGUIENTE · PRODUCTO OBJETIVO", 7.05, 2.48, 5.0, 0.3, 11, CORAL, True)
textbox(slide, "• Conectar el registro al KYC mock\n• Definir estados y acciones siguientes\n• Persistir la cuenta de demostración\n• Medir tiempo de onboarding con usuarios",
        7.05, 3.08, 4.9, 1.95, 18, INK)
textbox(slide, "El objetivo 8 → 3 min aún debe validarse.", 7.05, 5.05, 4.9, 0.36, 12, MUTED, True)
footer(slide, 6)

# 7. Outcome and next steps
slide = new_slide(DARK)
header(slide, "Éxito y próximos pasos", "Una meta medible; una ruta para validarla",
       "El objetivo de tiempo es una hipótesis de producto, aún no un resultado demostrado.", dark=True)
panel(slide, 0.82, 2.28, 4.15, 2.55, GREEN, GREEN)
textbox(slide, "MEDIANA DE ONBOARDING", 1.12, 2.62, 3.5, 0.3, 11, LIME, True)
textbox(slide, "8 → 3", 1.12, 3.15, 3.3, 0.85, 42, WHITE, True)
textbox(slide, "minutos · objetivo por validar con usuarios", 1.14, 4.12, 3.4, 0.42, 12, "DCE9E1")
textbox(slide, "PARA AVANZAR", 5.55, 2.35, 5.8, 0.28, 11, LIME, True)
next_steps = [
    "Acordar campos y criterios de aceptación W1.",
    "Conectar el formulario Angular al KYC mock.",
    "Definir estados y siguiente paso del usuario.",
    "Medir el tiempo real de finalización con usuarios.",
]
for i, item in enumerate(next_steps):
    y = 2.9 + i * 0.68
    textbox(slide, f"0{i + 1}", 5.55, y, 0.48, 0.28, 11, LIME, True)
    textbox(slide, item, 6.15, y - 0.02, 5.8, 0.42, 15, WHITE)
panel(slide, 5.55, 5.95, 6.7, 0.58, LIME, LIME)
textbox(slide, "SIGUIENTE DECISIÓN  ·  Validar datos KYC y estados del flujo", 5.82, 6.11, 6.15, 0.23, 11, DARK, True)
footer(slide, 7, True)

prs.core_properties.title = "Wallet para pagos cotidianos"
prs.core_properties.subject = "Propuesta de producto y demo del MVP de wallet en Perú"
prs.core_properties.author = "Equipo Wallet MVP"
prs.core_properties.keywords = "wallet, onboarding, KYC, QR, Perú, MVP"
prs.save(OUTPUT)
print(OUTPUT)
